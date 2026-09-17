import os, shutil, imaplib, email
from pathlib import Path
from email.header import decode_header
from datetime import datetime
from openpyxl import load_workbook

def env(name):
    value = os.getenv(name, '').strip()
    if not value: raise RuntimeError(f'Falta la variable {name}')
    return value

def decode_name(value):
    return ''.join(part.decode(enc or 'utf-8', errors='replace') if isinstance(part, bytes) else part for part, enc in decode_header(value or ''))

def main():
    user, password, sender = env('GMAIL_USER'), env('GMAIL_APP_PASSWORD'), env('FROM_SENDER')
    template = Path(env('TEMPLATE_XLSM')); output_dir = Path(os.getenv('OUTPUT_FOLDER', '/data/out'))
    if not template.exists(): raise RuntimeError(f'No existe la plantilla: {template}')
    mail = imaplib.IMAP4_SSL(os.getenv('IMAP_SERVER', 'imap.gmail.com'))
    try:
        mail.login(user, password); mail.select(os.getenv('IMAP_FOLDER', 'INBOX'))
        status, data = mail.search(None, f'(UNSEEN FROM "{sender}")')
        if status != 'OK' or not data[0]: print('No hay mensajes para procesar.'); return
        msg_id = data[0].split()[-1]; status, raw = mail.fetch(msg_id, '(RFC822)')
        if status != 'OK': raise RuntimeError('No se pudo leer el mensaje')
        msg = email.message_from_bytes(raw[0][1]); attachment = None
        for part in msg.walk():
            name = decode_name(part.get_filename())
            if part.get_content_disposition() == 'attachment' and name.lower().endswith('.xlsx'):
                dest = Path(os.getenv('DEST_FOLDER', '/data')) / Path(name).name; dest.parent.mkdir(parents=True, exist_ok=True)
                dest.write_bytes(part.get_payload(decode=True)); attachment = dest; break
        if not attachment: print('El mensaje no contiene un adjunto .xlsx.'); return
        wb = load_workbook(attachment, data_only=True)
        source = os.getenv('SOURCE_SHEET', 'Datos_conductor_y_vehiculo')
        target = os.getenv('DEST_SHEET', 'QUINCENA')
        if source not in wb.sheetnames: raise RuntimeError(f'No existe la hoja origen: {source}')
        rows = [[wb[source].cell(r,c).value for c in (1,2,5,8,3)] for r in range(2, wb[source].max_row+1)]
        rows = [row for row in rows if any(v not in (None, '') for v in row)]
        output_dir.mkdir(parents=True, exist_ok=True); out = output_dir / f'QUINCENA_{datetime.now():%Y%m%d_%H%M%S}.xlsm'; shutil.copy2(template, out)
        result = load_workbook(out, keep_vba=True)
        if target not in result.sheetnames: raise RuntimeError(f'No existe la hoja destino: {target}')
        ws = result[target]
        for i,row in enumerate(rows, 13):
            for j,value in enumerate(row, 1): ws.cell(i,j).value = value
        result.save(out); mail.store(msg_id, '+FLAGS', '\\Seen'); print(f'Salida generada: {out}')
    finally:
        try: mail.logout()
        except Exception: pass
if __name__ == '__main__': main()
